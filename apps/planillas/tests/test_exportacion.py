"""CU-05.5 Exportar planilla y boletas a PDF y Excel (RF-16, RNF-17)."""

import io

import openpyxl
import pytest
from django.urls import reverse

from apps.planillas.fachada import FachadaPlanilla
from apps.planillas.models import BoletaPago

from .utiles import ANIO, MES, contrato_con_marcacion, encargado, haber

pytestmark = pytest.mark.django_db


@pytest.fixture
def cerrada(client):
    haber("3000")
    contrato_con_marcacion(30)
    fachada = FachadaPlanilla()
    fachada.aprobar_y_cerrar(fachada.generar_preplanilla(ANIO, MES).planilla, encargado())
    client.force_login(encargado())
    return client


def test_planilla_en_excel(cerrada):
    respuesta = cerrada.get(reverse("planillas:exportar", args=[ANIO, MES, "xlsx"]))
    assert respuesta.status_code == 200
    hoja = openpyxl.load_workbook(io.BytesIO(respuesta.content)).active
    filas = list(hoja.iter_rows(min_row=3, values_only=True))
    assert filas[0][0] == "CI"
    assert filas[1][-1] == 3000.0


def test_boleta_en_excel(cerrada):
    nro = BoletaPago.objects.get().nro
    respuesta = cerrada.get(reverse("planillas:exportar_boleta", args=[nro, "xlsx"]))
    hoja = openpyxl.load_workbook(io.BytesIO(respuesta.content)).active
    conceptos = [fila[0] for fila in hoja.iter_rows(min_row=4, values_only=True)]
    assert "SUELDO GANADO" in conceptos
    assert "LÍQUIDO PAGABLE" in conceptos


def test_pdf_de_planilla_y_boleta(cerrada):
    nro = BoletaPago.objects.get().nro
    for url in (
        reverse("planillas:exportar", args=[ANIO, MES, "pdf"]),
        reverse("planillas:exportar_boleta", args=[nro, "pdf"]),
    ):
        respuesta = cerrada.get(url)
        assert respuesta["Content-Type"] == "application/pdf"
        assert respuesta.content.startswith(b"%PDF")


def test_formato_desconocido_y_planilla_sin_cerrar(cerrada):
    assert cerrada.get(reverse("planillas:exportar", args=[ANIO, MES, "doc"])).status_code == 404
    assert cerrada.get(reverse("planillas:exportar", args=[ANIO, 7, "pdf"])).status_code == 404
